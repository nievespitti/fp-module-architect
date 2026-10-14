#!/usr/bin/env python3
"""Genera el Cuaderno del Profesor / Diario de Aula (Excel) desde module.json.

Uso: python build_cuaderno.py <carpeta_modulo> [--out Cuaderno_Profesor_....xlsx]

Una fila por sesión REAL (horario semanal × calendario − no lectivos − periodos sin clase − FEOE).
Cada SdA aporta sus sesiones en sda["sesiones"]: lista, en orden, con una entrada por sesión:
  {"plan": "texto", "ce": "1.a, 1.b", "evidencia": "EV1"}   (ce y evidencia opcionales)
Si faltan entradas, la fila sale marcada "PENDIENTE DE PLANIFICAR" (el validador lo trata como error).
Hojas: Resumen (horas reales vs nominales, fórmulas) + una hoja por evaluación + Hitos.
"""
import os, sys, datetime as dt
sys.path.insert(0, os.path.dirname(__file__))
from common import load, d, fmt, sessions, blocked_days, DIAS
import openpyxl
from cuaderno_extra import add_extra_sheets
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def fill(c): return PatternFill(start_color=c, end_color=c, fill_type="solid")
HEADER_FILL, HEADER_FONT = fill("1F4E79"), Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1F4E79")
SUB_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
ZEBRA, WHITE = fill("F2F5F9"), fill("FFFFFF")
STYLE = {  # tipo: (relleno, fuente)
    "EVALUACION": (fill("D9E1F2"), Font(name="Calibri", size=10, bold=True, color="1F4E79")),
    "FESTIVO": (fill("FCE4D6"), Font(name="Calibri", size=10, bold=True, color="C65911")),
    "ENTREGA": (fill("FFF2CC"), Font(name="Calibri", size=10, bold=True, color="806000")),
    "FEOE": (fill("E2EFDA"), Font(name="Calibri", size=10, bold=True, color="375623")),
    "PENDIENTE": (fill("F8CBAD"), Font(name="Calibri", size=10, bold=True, color="C00000")),
}
B = Side(style="thin", color="D9D9D9")
BORDER = Border(left=B, right=B, top=B, bottom=B)
COLS = [("Sesión", 8, "center"), ("Fecha", 12, "center"), ("Día", 11, "center"), ("Horario", 13, "center"),
        ("Horas", 7, "center"), ("SdA / Unidad", 24, "left"), ("RA / CE", 14, "center"),
        ("Planificación / Contenido previsto", 46, "left"),
        ("Diario de aula / Trabajo realizado (registro del profesor)", 40, "left"),
        ("Incidencias / Tareas / Observaciones", 30, "left"), ("Evidencia Alexia", 16, "center")]

def sda_of(m, date):
    for s in m["sdas"]:
        if d(s["inicio"]) <= date <= d(s["fin"]):
            return s
    ui = m.get("unidad_inicial")
    if ui and d(ui["inicio"]) <= date <= d(ui["fin"]):
        return {"n": "UI", "titulo": ui["titulo"], "ra_principal": None, "sesiones": ui.get("sesiones", [])}
    return None

def build_rows(m):
    cal = m["calendario"]
    rows = []  # (fecha, orden, dict)
    counters = {}
    for ses in sessions(m):
        s = sda_of(m, ses["fecha"])
        if s is None:
            rows.append((ses["fecha"], ses["inicio"], dict(tipo="PENDIENTE", sda="SIN SdA ASIGNADA", ra="-", plan="PENDIENTE DE PLANIFICAR: sesión fuera de toda SdA",
                                                          ev="-", ses=ses)))
            continue
        i = counters.get(s["n"], 0)
        counters[s["n"]] = i + 1
        plan = s.get("sesiones", [])
        sda_label = s["titulo"] if s["n"] == "UI" else f"SdA{s['n']} · {s['titulo']}"
        ra_label = f"RA{s['ra_principal']}" if s["ra_principal"] is not None else "Evaluación inicial"
        if i < len(plan):
            p = plan[i] if isinstance(plan[i], dict) else {"plan": plan[i]}
            tipo = "ENTREGA" if p.get("evidencia") else "NORMAL"
            rows.append((ses["fecha"], ses["inicio"], dict(tipo=tipo, sda=sda_label, ra=p.get("ce") or ra_label,
                                                          plan=p["plan"], ev=p.get("evidencia", "-"), ses=ses)))
        else:
            rows.append((ses["fecha"], ses["inicio"], dict(tipo="PENDIENTE", sda=sda_label, ra=ra_label,
                                                          plan="PENDIENTE DE PLANIFICAR", ev="-", ses=ses)))
    # hitos de calendario (evaluaciones, entregas de notas, actos oficiales, festivos dentro del periodo)
    for e in cal["evaluaciones"]:
        rows.append((d(e["fecha"]), "00:00", dict(tipo="EVALUACION", sda=e.get("nombre", f"{e['n']}ª EVALUACIÓN"), ra="-",
                                                  plan=e.get("descripcion", "Sesión oficial de evaluación. Cierre de notas en Alexia."),
                                                  ev=e.get("evidencia", "Actas"), ses=None)))
    for h in cal.get("hitos", []):
        rows.append((d(h["fecha"]), "00:00", dict(tipo=h.get("tipo", "EVALUACION"), sda=h["nombre"], ra="-", plan=h.get("descripcion", h["nombre"]), ev="-", ses=None)))
    for n in cal.get("no_lectivos", []):
        dd = d(n["fecha"])
        if d(cal["inicio_clases"]) <= dd <= d(cal.get("fin_curso", cal["fin_clases_modulo"])) and dd.weekday() < 5 and any(h["dia"] == DIAS[dd.weekday()] for h in m["horario"]):
            rows.append((dd, "00:00", dict(tipo="FESTIVO", sda="DÍA NO LECTIVO", ra="-", plan=n.get("motivo", "No lectivo"), ev="-", ses=None)))
    for p in cal.get("periodos_sin_clase", []):
        rows.append((d(p["inicio"]), "00:00", dict(tipo="FESTIVO", sda=p.get("motivo", "Sin clase").upper(), ra="-",
                                                  plan=f"{p.get('motivo','Sin clase')}: del {fmt(d(p['inicio']))} al {fmt(d(p['fin']))}", ev="-", ses=None)))
    if cal.get("feoe"):
        f = cal["feoe"]
        rows.append((d(f["inicio"]), "00:00", dict(tipo="FEOE", sda="FEOE · INICIO", ra="-", plan=f"Inicio del periodo FEOE (hasta {fmt(d(f['fin']))}). Tutorías de seguimiento: {f.get('tutorias','según centro')}.", ev="Ficha FEOE", ses=None)))
        a, e = d(f["inicio"]), d(f["fin"])
        wd = {"lunes": 0, "martes": 1, "miércoles": 2, "miercoles": 2, "jueves": 3, "viernes": 4}.get(str(f.get("tutorias", "")).lower())
        blk = blocked_days({"periodos_sin_clase": cal.get("periodos_sin_clase", []), "no_lectivos": cal.get("no_lectivos", [])})
        fechas = [d(x) for x in f.get("tutorias_fechas", [])]
        for cur in fechas:
            rows.append((cur, "08:00", dict(tipo="FEOE", sda="FEOE · Tutoría de seguimiento", ra="CE dualizados",
                                            plan=f"Tutoría de seguimiento en el centro ({f.get('tutorias', 'según calendario')}): revisar diario EVAGD, contrastar tareas de empresa con los CE dualizados y detectar CE sin cubrir (activar alternativa en el centro). Se aprovecha para recuperaciones pendientes.",
                                            ev="Diario EVAGD", ses=None)))
        if wd is not None and not fechas:
            cur = a
            while cur <= e:
                if cur.weekday() == wd and cur not in blk and cur != a:
                    rows.append((cur, "08:00", dict(tipo="FEOE", sda="FEOE · Tutoría de seguimiento", ra="CE dualizados",
                                                    plan="Tutoría en el centro: revisar diario EVAGD, contrastar tareas de empresa con los CE dualizados y detectar CE sin cubrir (activar alternativa en el centro).",
                                                    ev="Diario EVAGD", ses=None)))
                cur += dt.timedelta(days=1)
        rows.append((d(f["fin"]), "99:99", dict(tipo="FEOE", sda="FEOE · FIN", ra="-", plan="Fin del periodo FEOE. Recogida de valoraciones y firmas.", ev="Acta FEOE", ses=None)))
    rows.sort(key=lambda r: (r[0], r[1]))
    return rows

def sheet(wb, name, title, sub, rows, start_no):
    ws = wb.create_sheet(name)
    n = len(COLS)
    for r, (txt, font) in enumerate([(title, TITLE_FONT), (sub, SUB_FONT)], 1):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=n)
        ws.cell(r, 1, txt).font = font
    ws.row_dimensions[1].height = 24
    for i, (nm, w, _) in enumerate(COLS, 1):
        c = ws.cell(4, i, nm)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[4].height = 30
    no, r = start_no, 5
    for fecha, _, x in rows:
        ses = x["ses"]
        vals = [no if ses else "", fecha.strftime("%d/%m/%Y"), DIAS[fecha.weekday()],
                f"{ses['inicio']}-{ses['fin']}" if ses else "", ses["horas"] if ses else 0,
                x["sda"], x["ra"], x["plan"], "", "", x["ev"]]
        if ses: no += 1
        fl, ft = STYLE.get(x["tipo"], (ZEBRA if r % 2 == 0 else WHITE, Font(name="Calibri", size=10)))
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.fill, c.border = fl, BORDER
            c.font = ft if x["tipo"] in STYLE and i in (1, 2, 6, 8, 11) else Font(name="Calibri", size=10)
            c.alignment = Alignment(horizontal=COLS[i - 1][2], vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = 30 if len(str(x["plan"])) < 60 else 44
        r += 1
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(n)}{r - 1}"
    return no, r - 1

def resumen(wb, m, sheets):
    ws = wb.create_sheet("Resumen", 0)
    mod = m["modulo"]
    ws["A1"] = f"{mod['nombre']} ({mod['codigo']}) · {mod['curso']} · {mod['centro']} · {mod['anio']}"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = f"Docente: {m['docente']}. Horario: " + "; ".join(f"{h['dia']} {h['inicio']}-{h['fin']}" for h in m["horario"])
    ws["A2"].font = SUB_FONT
    head = ["SdA", "RA principal", "Inicio", "Fin", "Horas declaradas", "Horas en cuaderno (fórmula)"]
    for i, h in enumerate(head, 1):
        c = ws.cell(4, i, h); c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    r = 5
    for s in m["sdas"]:
        ws.cell(r, 1, f"SdA{s['n']} · {s['titulo']}"); ws.cell(r, 2, f"RA{s['ra_principal']}")
        ws.cell(r, 3, fmt(d(s["inicio"]))); ws.cell(r, 4, fmt(d(s["fin"]))); ws.cell(r, 5, s["horas"])
        refs = "+".join(f"SUMIF('{sn}'!F:F,\"SdA{s['n']} · {s['titulo']}\",'{sn}'!E:E)" for sn in sheets)
        ws.cell(r, 6, "=" + refs)
        for i in range(1, 7): ws.cell(r, i).border = BORDER
        r += 1
    ws.cell(r, 1, "TOTAL").font = Font(bold=True)
    ws.cell(r, 5, f"=SUM(E5:E{r-1})"); ws.cell(r, 6, f"=SUM(F5:F{r-1})")
    ws.cell(r + 1, 1, "Horas nominales del módulo (norma)"); ws.cell(r + 1, 5, mod["horas_nominales"])
    ws.cell(r + 2, 1, "Horas reales de clase (horario × calendario)"); ws.cell(r + 2, 5, sum(x["horas"] for x in sessions(m)))
    ws.cell(r + 3, 1, "Diferencia real − nominal"); ws.cell(r + 3, 5, f"=E{r+2}-E{r+1}")
    dual = m.get("dual")
    if dual:
        ws.cell(r + 5, 1, "Horas dualizadas en FEOE"); ws.cell(r + 5, 5, dual["horas"])
        ws.cell(r + 6, 1, "Aula + FEOE (debe igualar las nominales)"); ws.cell(r + 6, 5, f"=E{r}+E{r+5}")
        ws.cell(r + 8, 1, "CE dualizados en FEOE").font = Font(bold=True)
        for i, x in enumerate(dual["ces"]):
            ws.cell(r + 9 + i, 1, f"CE {x['ce']}"); ws.cell(r + 9 + i, 5, x["horas"]); ws.cell(r + 9 + i, 6, x["evidencia_empresa"][:120])
    ws.cell(r + 4, 1, "Si la diferencia es negativa, la programación debe justificar el reparto (docente decide).").font = Font(italic=True, color="C00000")
    for i, w in enumerate([44, 14, 12, 12, 18, 26], 1): ws.column_dimensions[get_column_letter(i)].width = w

def main():
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2: sys.exit(__doc__)
    root = sys.argv[1]
    m = load(os.path.join(root, "module.json"))
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else os.path.join(
        root, f"Cuaderno_Profesor_Diario_Aula_{m['modulo']['nombre'].replace(' ', '_')}_{m['modulo']['anio']}.xlsx")
    rows = build_rows(m)
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    evs = sorted(m["calendario"]["evaluaciones"], key=lambda e: e["fecha"])
    bounds, prev = [], dt.date.min
    for e in evs:
        bounds.append((e["n"], prev, d(e["fecha"]))); prev = d(e["fecha"]) + dt.timedelta(days=1)
    if evs and "final" in evs[-1].get("nombre", "").lower():
        bounds[-1] = (bounds[-1][0], bounds[-1][1], dt.date.max)   # el cierre de curso va dentro de la evaluación final
    else:
        bounds.append((len(evs) + 1, prev, dt.date.max))
    names, no = [], 1
    for n, a, b in bounds:
        part = [r for r in rows if a <= r[0] <= b]
        if not part: continue
        enombre = next((e.get("nombre", "") for e in evs if e["n"] == n), "")
        nm = "Evaluación final" if "final" in enombre.lower() else (f"{n}ª Evaluación" if n <= len(evs) else "Cierre de curso")
        sub = f"{m['modulo']['nombre']} · {m['modulo']['curso']} · " + " · ".join(sorted({x['sda'] for _, _, x in part if x['tipo'] in ('NORMAL', 'ENTREGA')}))
        no, _ = sheet(wb, nm, f"Cuaderno del Profesor · Diario de Aula · {nm} ({m['modulo']['anio']})", sub[:250], part, no)
        names.append(nm)
    resumen(wb, m, names)
    add_extra_sheets(wb, m)
    wb.save(out)
    pend = sum(1 for r in rows if r[2]["tipo"] == "PENDIENTE")
    print(f"Guardado: {out}\nSesiones reales: {sum(1 for r in rows if r[2]['ses'])}. Filas PENDIENTES de planificar: {pend}")

if __name__ == "__main__":
    main()
