#!/usr/bin/env python3
"""Validador de auditoría. Genera Informe_Validacion.md. NUNCA lo escribe el generador de contenido.

Uso: python validate.py <carpeta_modulo> [--out Informe_Validacion.md]
Espera <carpeta_modulo>/module.json y <carpeta_modulo>/ra_oficiales.json (salida de extract_official_ce.py).
Código de salida 1 si hay ERRORES.
"""
import glob, os, re, sys, datetime as dt
sys.path.insert(0, os.path.dirname(__file__))
from common import load, d, fmt, sessions, all_ce, blocked_days

ERR, WARN, OK = [], [], []
def err(m): ERR.append(m)
def warn(m): WARN.append(m)
def ok(m): OK.append(m)

TOL = 0.01
MIN_BYTES = {"ficha": 6000, "guia": 5000, "rubrica": 6000, "examen": 2500,
             "mapeo": 2500, "protocolo": 2000, "dua": 2000}
PLACEHOLDER = re.compile(r"Contenido SdA|lorem ipsum|\{\{|\}\}|\[X\]|\[COMPLETAR|\[INSERTAR|<<[A-Z_]+>>|Jortiz|\[\[REDACTAR", re.I)
PLACEHOLDER_CS = re.compile(r"TODO|TBD")
CE_REF = re.compile(r"\bCE\s?(\d+)\.([a-z])\b")
MARK = re.compile(r"<!--CE:(\d+\.[a-z])-->(.*?)<!--/CE-->", re.S)
norm = lambda s: re.sub(r"\s+", " ", s).strip()


def kind(fn):
    f = fn.lower()
    for k, keys in [("ficha", ["ficha"]), ("guia", ["guia", "guía"]), ("rubrica", ["rubrica", "rúbrica"]),
                    ("examen", ["examen"]), ("mapeo", ["mapeo"]), ("protocolo", ["protocolo"]),
                    ("dua", ["dua", "medidas"])]:
        if any(x in f for x in keys):
            return k
    return None


def check_structure(m):
    for k in ["modulo", "docente", "horario", "calendario", "ras", "sdas"]:
        if k not in m:
            err(f"module.json: falta la clave '{k}'")
    if ERR:
        return False
    for k in ["nombre", "codigo", "curso", "centro", "anio", "horas_nominales"]:
        if k not in m["modulo"]:
            err(f"modulo.{k} ausente")
    for k in ["inicio_clases", "fin_clases_modulo", "evaluaciones"]:
        if k not in m["calendario"]:
            err(f"calendario.{k} ausente")
    return not ERR


def check_official(m, off):
    if not off:
        warn("No hay ra_oficiales.json: los CE NO se pueden comprobar contra el texto oficial.")
        return
    n0 = len(ERR)
    if len(off["ras"]) != len(m["ras"]):
        err(f"Nº de RA: module.json={len(m['ras'])} oficial={len(off['ras'])}")
    for ora in off["ras"]:
        ra = next((r for r in m["ras"] if r["n"] == ora["n"]), None)
        if not ra:
            err(f"RA{ora['n']} oficial no está en module.json")
            continue
        ids_o = [c["id"] for c in ora["ces"]]
        ids_m = [c["id"] for c in ra["ces"]]
        if ids_o != ids_m:
            err(f"RA{ora['n']}: CE distintos. oficial={ids_o} module={ids_m}")
        for oc in ora["ces"]:
            mc = next((c for c in ra["ces"] if c["id"] == oc["id"]), None)
            if mc and norm(mc["texto"]) != norm(oc["texto"]):
                err(f"CE {ora['n']}.{oc['id']}: texto NO literal respecto al oficial")
    if len(ERR) == n0:
        ok(f"{sum(len(r['ces']) for r in m['ras'])} CE literales = texto oficial ({len(m['ras'])} RA)")


def check_weights(m):
    n0 = len(ERR)
    s = sum(r["peso_ra"] for r in m["ras"])
    if abs(s - 100) >= TOL:
        err(f"Pesos RA suman {s:g} (debe ser 100)")
    for ra in m["ras"]:
        s = sum(c["peso"] for c in ra["ces"])
        if abs(s - 100) >= TOL:
            err(f"RA{ra['n']}: pesos de CE suman {s:g} (debe ser 100)")
    for sda in m["sdas"]:
        s = sum(e["peso"] for e in sda["evidencias"])
        if abs(s - 100) >= TOL:
            err(f"SdA{sda['n']}: pesos de evidencias suman {s:g} (debe ser 100)")
        ind = [e for e in sda["evidencias"] if e.get("tipo") == "examen_individual"]
        if not ind:
            err(f"SdA{sda['n']}: no hay evidencia tipo 'examen_individual' (presencial, sin internet)")
        elif not any(20 <= e["peso"] <= 35 for e in ind):
            warn(f"SdA{sda['n']}: examen individual con peso fuera de 20-35%")
        if len(sda["evidencias"]) < 4:
            warn(f"SdA{sda['n']}: solo {len(sda['evidencias'])} evidencias")
    if len(ERR) == n0:
        ok("Pesos: RA=100, CE por RA=100, evidencias por SdA=100")


def check_coverage(m):
    n0 = len(ERR)
    ces = all_ce(m)
    used = set()
    for sda in m["sdas"]:
        if not any(r["n"] == sda["ra_principal"] for r in m["ras"]):
            err(f"SdA{sda['n']}: ra_principal inexistente")
        for ev in sda["evidencias"]:
            for c in ev["ces"]:
                if c not in ces:
                    err(f"SdA{sda['n']} {ev['id']}: CE '{c}' no existe en el módulo")
                used.add(c)
    dual = {x["ce"] for x in m.get("dual", {}).get("ces", [])}
    miss = [c for c in ces if c not in used and c not in dual]
    if miss:
        err("CE no evaluados por ninguna evidencia: " + ", ".join(miss))
    for ra in m["ras"]:
        sd = [s for s in m["sdas"] if s["ra_principal"] == ra["n"] or ra["n"] in s.get("ras_complementarios", [])]
        if not sd:
            err(f"RA{ra['n']} no se trabaja en ninguna SdA")
        if ra.get("transversal") and len(sd) < 2:
            warn(f"RA{ra['n']} es transversal pero aparece en {len(sd)} SdA")
        if not ra.get("evaluaciones"):
            warn(f"RA{ra['n']}: sin mapeo a evaluaciones")
    if len(ERR) == n0:
        ok("Todos los CE están evaluados; todos los RA se trabajan en alguna SdA")


def check_hours(m):
    n0 = len(ERR)
    cal = m["calendario"]
    nom = m["modulo"]["horas_nominales"]
    real = sum(s["horas"] for s in sessions(m))
    msg = f"Horas reales de clase (horario × calendario) = {real} h; nominales = {nom} h"
    if real < nom * 0.95:
        warn(msg + f". Déficit {nom - real} h que debe cubrirse con horas dualizadas en FEOE.")
    else:
        ok(msg)
    dual = m.get("dual")
    dh = dual["horas"] if dual else 0
    if dual:
        ces = all_ce(m)
        sh = sum(x["horas"] for x in dual["ces"])
        if sh != dh:
            err(f"FEOE dual: horas por CE suman {sh} y el total declarado es {dh}")
        for x in dual["ces"]:
            if x["ce"] not in ces:
                err(f"FEOE dual: CE {x['ce']} no existe")
            for k in ("evidencia_empresa", "alternativa_centro"):
                if not x.get(k):
                    err(f"FEOE dual: CE {x['ce']} sin '{k}'")
        if not cal.get("feoe"):
            err("Hay horas dualizadas pero no hay periodo FEOE en el calendario")
        ok(f"FEOE dual: {dh} h repartidas en {len(dual['ces'])} CE, cada uno con evidencia de empresa y alternativa en el centro")
    elif real < nom * 0.95:
        warn("Déficit de horas sin dualización declarada (module.json 'dual')")
    ui = m.get("unidad_inicial")
    aula = sum(s["horas"] for s in m["sdas"]) + (ui["horas"] if ui else 0)
    if aula + dh != nom:
        err(f"Horas de aula asignadas ({aula}) + FEOE dual ({dh}) = {aula + dh}; deben ser las nominales ({nom})")
    else:
        ok(f"Aula {aula} h + FEOE dual {dh} h = {nom} h nominales")
    tot = aula
    if tot != real:
        warn(f"Horas asignadas a SdA = {tot} h; horas reales disponibles = {real} h")
    else:
        ok(f"Horas asignadas a SdA = horas reales = {real} h")
    prev_end = d(ui["fin"]) if ui else None
    for s in sorted(m["sdas"], key=lambda x: x["inicio"]):
        a, b = d(s["inicio"]), d(s["fin"])
        if prev_end and a <= prev_end:
            warn(f"SdA{s['n']} solapa con la anterior ({s['inicio']})")
        prev_end = b
        r = sum(x["horas"] for x in sessions(m, s["inicio"], s["fin"]))
        if r != s["horas"]:
            warn(f"SdA{s['n']}: declara {s['horas']} h y el calendario da {r} h entre {s['inicio']} y {s['fin']}")
        if b > d(cal["fin_clases_modulo"]):
            err(f"SdA{s['n']} termina {s['fin']} después del fin de clases del módulo ({cal['fin_clases_modulo']})")
        if cal.get("feoe") and a <= d(cal["feoe"]["fin"]) and b >= d(cal["feoe"]["inicio"]):
            err(f"SdA{s['n']} invade la FEOE")
        evs = s.get("evaluaciones", [])
        if len(evs) == 1:
            ev = next((x for x in cal["evaluaciones"] if x["n"] == evs[0]), None)
            if ev and b > d(ev["fecha"]):
                err(f"SdA{s['n']} termina {s['fin']} después de la evaluación {evs[0]} ({ev['fecha']})")
    if len(ERR) == n0:
        ok("Ninguna SdA invade FEOE, sobrepasa fin de clases ni cruza su evaluación")


def check_files(m, root):
    n0 = len(ERR)
    ces = all_ce(m)
    for sda in m["sdas"]:
        n = sda["n"]
        dirs = glob.glob(os.path.join(root, f"SdA{n}_*")) + glob.glob(os.path.join(root, f"{n:02d}_SdA*"))
        dirs = sorted({x for x in dirs if os.path.isdir(x)})
        if not dirs:
            err(f"SdA{n}: carpeta no encontrada (SdA{n}_*)")
            continue
        if len(dirs) > 1:
            err(f"SdA{n}: carpetas duplicadas: {[os.path.basename(x) for x in dirs]}")
        folder = dirs[0]
        found = {}
        for f in os.listdir(folder):
            k = kind(f)
            if k and f.lower().endswith((".md", ".docx")):
                found.setdefault(k, []).append(f)
        allowed = set()
        for ev in sda["evidencias"]:
            allowed.update(ev["ces"])
        for k, mn in MIN_BYTES.items():
            if k not in found:
                err(f"SdA{n}: falta documento '{k}'")
                continue
            for f in found[k]:
                p = os.path.join(folder, f)
                sz = os.path.getsize(p)
                if sz < mn:
                    err(f"SdA{n}/{f}: {sz} B < mínimo {mn} B (¿esqueleto?)")
                if not f.endswith(".md"):
                    continue
                txt = open(p, encoding="utf-8", errors="replace").read()
                mm = PLACEHOLDER.search(txt) or PLACEHOLDER_CS.search(txt)
                if mm:
                    err(f"SdA{n}/{f}: placeholder/basura '{mm.group(0)}'")
                refs = {f"{a}.{b}" for a, b in CE_REF.findall(txt)} | {c for c, _ in MARK.findall(txt)}
                bad = sorted(r for r in refs if r not in allowed)
                if bad:
                    err(f"SdA{n}/{f}: cita CE que no pertenecen a esta SdA: {bad}")
                if k in ("rubrica", "mapeo"):
                    absent = sorted(c for c in allowed if c not in refs)
                    if absent:
                        err(f"SdA{n}/{f}: no menciona los CE {absent}")
                for cid, body in MARK.findall(txt):
                    if cid in ces and norm(body) != norm(ces[cid][1]["texto"]):
                        err(f"SdA{n}/{f}: CE {cid} marcado con texto no literal")
    if len(ERR) == n0:
        ok("Documentos por SdA completos, sin placeholders, CE coherentes con su SdA")


def check_module_docs(m, root):
    def has(pat):
        return glob.glob(os.path.join(root, "**", pat), recursive=True)
    for label, pats, mn in [("Programación didáctica", ["Programacion*.docx", "Programacion*.md"], 20000),
                            ("Tabla maestra", ["Tabla_Maestra*.xlsx"], 5000),
                            ("Cuaderno/temporalización (Excel)", ["Cuaderno*.xlsx"], 12000)]:
        fs = [f for p in pats for f in has(p)]
        if not fs:
            err(f"Falta: {label}")
        elif max(os.path.getsize(f) for f in fs) < mn:
            err(f"{label}: archivo sospechosamente pequeño (<{mn} B)")
        else:
            ok(f"{label}: presente ({os.path.basename(fs[0])})")
            if label.startswith("Cuaderno"):
                import openpyxl
                wb = openpyxl.load_workbook(fs[0])
                pend = sum(1 for ws in wb for row in ws.iter_rows(values_only=True) for v in row if isinstance(v, str) and "PENDIENTE" in v)
                if pend:
                    err(f"Cuaderno Excel: {pend} celda(s) 'PENDIENTE' (sesiones sin planificar)")
    if m["calendario"].get("feoe") and not has("*FEOE*"):
        err("Hay FEOE y falta la ficha de seguimiento FEOE")
    if not has("Memoria*"):
        warn("Falta Memoria final de práctica docente")


def report(m, out):
    head = "NO APTO — %d error(es)" % len(ERR) if ERR else "APTO (sin errores)"
    L = [f"# Informe de validación — {m['modulo']['nombre']} ({m['modulo']['codigo']})",
         f"_Generado por `validate.py` el {dt.date.today().isoformat()}. No editar a mano._", "",
         f"**Resultado: {head}** · {len(WARN)} aviso(s)", ""]
    L += ["## Errores"] + ([f"- ❌ {x}" for x in ERR] or ["- ninguno"])
    L += ["", "## Avisos"] + ([f"- ⚠️ {x}" for x in WARN] or ["- ninguno"])
    L += ["", "## Comprobaciones superadas"] + [f"- ✅ {x}" for x in OK]
    open(out, "w", encoding="utf-8").write("\n".join(L) + "\n")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    root = sys.argv[1]
    out = sys.argv[sys.argv.index("--out") + 1] if "--out" in sys.argv else os.path.join(root, "Informe_Validacion.md")
    m = load(os.path.join(root, "module.json"))
    offp = os.path.join(root, "ra_oficiales.json")
    off = load(offp) if os.path.exists(offp) else None
    try:
        if check_structure(m):
            check_official(m, off)
            check_weights(m)
            check_coverage(m)
            check_hours(m)
            check_files(m, root)
            check_module_docs(m, root)
    except (ValueError, KeyError) as e:
        sys.exit(f"module.json incompleto o con un dato no válido ({type(e).__name__}: {e}). "
                 "Rellena horario, calendario (fechas ISO aaaa-mm-dd) y sdas; ver references/module_schema.md "
                 "y examples/modulo-ejemplo/module.json.")
    report(m, out)
    print(open(out, encoding="utf-8").read())
    sys.exit(1 if ERR else 0)
