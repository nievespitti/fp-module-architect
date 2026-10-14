"""Hojas de gestión del cuaderno del profesor: alumnado, asistencia, actitud, FEOE y calificaciones con fórmulas.

Las fórmulas usan los pesos de module.json (evidencias por SdA, CE por RA, RA del módulo), así que las notas
de CE, RA y módulo se calculan solas al introducir las notas de las evidencias. No contiene datos personales:
el docente escribe los nombres en la hoja «Alumnado».
"""
import datetime as dt
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from common import d, sessions, all_ce

def fill(c): return PatternFill(start_color=c, end_color=c, fill_type="solid")
HF, HFONT = fill("1F4E79"), Font(name="Calibri", size=10, bold=True, color="FFFFFF")
INPUT, CALC, SUB = fill("FFF9DB"), fill("E8F1FA"), fill("D9E1F2")
S = Side(style="thin", color="BFBFBF")
BD = Border(left=S, right=S, top=S, bottom=S)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
DEFAULT_ACTITUD = ["Llega a la hora y permanece en su puesto", "Mantiene el móvil guardado salvo que la actividad lo pida",
                   "Muestra predisposición positiva ante la tarea", "Trabaja en equipo: escucha, reparte, cumple su parte",
                   "Cuida el material y respeta la confidencialidad", "Responde con respeto al profesor y a sus compañeros"]

def hdr(ws, row, col, text, width=None, fillc=HF, font=HFONT):
    c = ws.cell(row, col, text)
    c.fill, c.font, c.border, c.alignment = fillc, font, BD, CENTER
    if width: ws.column_dimensions[L(col)].width = width
    return c

def title(ws, text, sub=None):
    ws["A1"] = text
    ws["A1"].font = Font(name="Calibri", size=13, bold=True, color="1F4E79")
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="595959")

def n_alumnos(m):
    return int(m.get("grupo", {}).get("alumnos", 10))

def alumnado(wb, m):
    ws = wb.create_sheet("Alumnado")
    title(ws, "Alumnado del grupo", "Escriba aquí los nombres. Las demás hojas los toman de esta. No anote datos de necesidades de apoyo sin acreditación oficial.")
    for i, (t, w) in enumerate([("Nº", 6), ("Nombre y apellidos", 38), ("Observaciones (sin datos sin acreditar)", 50)], 1):
        hdr(ws, 4, i, t, w)
    for r in range(5, 5 + n_alumnos(m)):
        ws.cell(r, 1, r - 4).border = BD
        for c in (2, 3):
            ws.cell(r, c).fill, ws.cell(r, c).border = INPUT, BD
    return 5, 4 + n_alumnos(m)

def name_ref(r):
    return f"=IF(Alumnado!B{r}=\"\",\"Alumno/a {r-4}\",Alumnado!B{r})"

def asistencia(wb, m):
    ws = wb.create_sheet("Asistencia")
    title(ws, "Asistencia por sesión", "Códigos: F falta sin justificar · J falta justificada · R retraso. Casilla vacía = asistió.")
    ses = sessions(m)
    hdr(ws, 4, 1, "Alumno/a", 24)
    for j, s in enumerate(ses):
        col = 2 + j
        hdr(ws, 3, col, f"S{j+1}", 5, SUB, Font(name="Calibri", size=8, bold=True))
        c = hdr(ws, 4, col, s["fecha"].strftime("%d/%m"), 5, HF, Font(name="Calibri", size=8, bold=True, color="FFFFFF"))
    last = 1 + len(ses)
    sc = last + 1
    for k, t in enumerate(["Faltas sin justificar", "Faltas justificadas", "Retrasos", "% asistencia"]):
        hdr(ws, 4, sc + k, t, 12)
    dv = DataValidation(type="list", formula1='"F,J,R"', allow_blank=True)
    ws.add_data_validation(dv)
    for r in range(5, 5 + n_alumnos(m)):
        ws.cell(r, 1, name_ref(r)).border = BD
        for col in range(2, last + 1):
            ws.cell(r, col).fill, ws.cell(r, col).border = INPUT, BD
            dv.add(ws.cell(r, col))
        rng = f"B{r}:{L(last)}{r}"
        ws.cell(r, sc, f'=COUNTIF({rng},"F")').fill = CALC
        ws.cell(r, sc + 1, f'=COUNTIF({rng},"J")').fill = CALC
        ws.cell(r, sc + 2, f'=COUNTIF({rng},"R")').fill = CALC
        pc = ws.cell(r, sc + 3, f'=1-({L(sc)}{r}+{L(sc+1)}{r})/{len(ses)}')
        pc.fill, pc.number_format = CALC, "0%"
        for k in range(4): ws.cell(r, sc + k).border = BD
    ws.freeze_panes = "B5"

def actitud(wb, m):
    items = m.get("actitud_items", DEFAULT_ACTITUD)
    ws = wb.create_sheet("Actitud")
    title(ws, "Registro de actitud profesional (1 a 4)", "4 siempre · 3 casi siempre · 2 a veces · 1 rara vez. No modifica la nota de los CE; se entrega al tutor de la FEOE.")
    evs = [e.get("nombre", f"{e['n']}.ª evaluación") for e in sorted(m["calendario"]["evaluaciones"], key=lambda e: e["fecha"])]
    hdr(ws, 4, 1, "Alumno/a", 24)
    col = 2
    dv = DataValidation(type="whole", operator="between", formula1="1", formula2="4", allow_blank=True)
    ws.add_data_validation(dv)
    blocks = []
    for ev in evs:
        start = col
        ws.merge_cells(start_row=3, start_column=col, end_row=3, end_column=col + len(items))
        hdr(ws, 3, col, ev, None, SUB, Font(name="Calibri", size=10, bold=True, color="1F4E79"))
        for k, it in enumerate(items):
            hdr(ws, 4, col, it, 15); col += 1
        hdr(ws, 4, col, "Media", 9); col += 1
        blocks.append((start, col - 2, col - 1))
    ws.row_dimensions[4].height = 75
    for r in range(5, 5 + n_alumnos(m)):
        ws.cell(r, 1, name_ref(r)).border = BD
        for a, b, mean in blocks:
            for c in range(a, b + 1):
                ws.cell(r, c).fill, ws.cell(r, c).border = INPUT, BD
                dv.add(ws.cell(r, c))
            cc = ws.cell(r, mean, f'=IF(COUNT({L(a)}{r}:{L(b)}{r})=0,"",AVERAGE({L(a)}{r}:{L(b)}{r}))')
            cc.fill, cc.border, cc.number_format = CALC, BD, "0.0"
    ws.freeze_panes = "B5"

def feoe(wb, m):
    dual = m.get("dual")
    if not dual: return {}
    ws = wb.create_sheet("FEOE seguimiento")
    f = m["calendario"].get("feoe", {})
    title(ws, "Seguimiento de los CE dualizados en la FEOE",
          f"FEOE del {d(f['inicio']).strftime('%d/%m/%Y')} al {d(f['end'] if 'end' in f else f['fin']).strftime('%d/%m/%Y')}. Por cada CE, anote la nota del tutor de empresa o, si la empresa no lo cubre, la de la alternativa en el centro (solo una de las dos).")
    for i, (t, w) in enumerate([("Alumno/a", 24), ("Empresa", 24), ("Tutor/a de empresa", 22), ("Horas dual cumplidas", 12)], 1):
        hdr(ws, 4, i, t, w)
    ces = all_ce(m)
    refs = {}
    col = 5
    for x in dual["ces"]:
        k = x["ce"]
        ws.merge_cells(start_row=2 + 1, start_column=col, end_row=3, end_column=col + 2)
        hdr(ws, 3, col, f"CE {k} · {x['horas']} h", None, SUB, Font(name="Calibri", size=10, bold=True, color="1F4E79"))
        hdr(ws, 4, col, "Tutor empresa 0-10", 11); hdr(ws, 4, col + 1, "Alternativa centro 0-10", 11); hdr(ws, 4, col + 2, "Nota CE", 9)
        refs[k] = col + 2
        col += 3
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="10", allow_blank=True)
    ws.add_data_validation(dv)
    ws.row_dimensions[4].height = 40
    for r in range(5, 5 + n_alumnos(m)):
        ws.cell(r, 1, name_ref(r)).border = BD
        for c in (2, 3, 4):
            ws.cell(r, c).fill, ws.cell(r, c).border = INPUT, BD
        for k, c3 in refs.items():
            for c in (c3 - 2, c3 - 1):
                ws.cell(r, c).fill, ws.cell(r, c).border = INPUT, BD
                dv.add(ws.cell(r, c))
            g = ws.cell(r, c3, f'=IF(ISNUMBER({L(c3-1)}{r}),{L(c3-1)}{r},IF(ISNUMBER({L(c3-2)}{r}),{L(c3-2)}{r},""))')
            g.fill, g.border, g.number_format = CALC, BD, "0.0"
    ws.freeze_panes = "E5"
    return refs

def calificaciones(wb, m, feoe_refs):
    ws = wb.create_sheet("Calificaciones")
    title(ws, "Calificaciones con cálculo automático",
          "Escriba solo las casillas amarillas (notas 0-10). CE, RA y nota del módulo se calculan con los pesos de la Tabla maestra. Una nota recuperada sustituye a la anterior. Las casillas calculadas quedan vacías hasta tener todos los datos.")
    n = n_alumnos(m)
    hdr(ws, 4, 1, "Alumno/a", 24)
    col = 2
    ev_cols = {}   # (sda, evid) -> col
    sda_blocks = []
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="10", allow_blank=True)
    ws.add_data_validation(dv)
    ws.cell(3, 1, "Peso →").font = Font(bold=True, size=9)
    for s in m["sdas"]:
        start = col
        for e in s["evidencias"]:
            hdr(ws, 4, col, f"SdA{s['n']} {e['id']}", 8)
            w = ws.cell(3, col, e["peso"]); w.fill, w.border, w.alignment = SUB, BD, CENTER
            ev_cols[(s["n"], e["id"])] = col
            col += 1
        sda_blocks.append((s["n"], start, col - 1))
    sda_nota = {}
    for sn, a, b in sda_blocks:
        hdr(ws, 4, col, f"Nota SdA{sn}", 9, SUB, Font(name="Calibri", size=10, bold=True, color="1F4E79"))
        sda_nota[sn] = col
        col += 1
    # CE columns
    ces = all_ce(m)
    dual = {x["ce"] for x in m.get("dual", {}).get("ces", [])}
    ce_cols = {}
    ra_ranges = {}
    for ra in m["ras"]:
        start = col
        for ce in ra["ces"]:
            k = f"{ra['n']}.{ce['id']}"
            hdr(ws, 4, col, f"CE {k}", 7)
            w = ws.cell(3, col, ce["peso"]); w.fill, w.border, w.alignment = SUB, BD, CENTER
            ce_cols[k] = col
            col += 1
        ra_ranges[ra["n"]] = (start, col - 1)
    ra_cols = {}
    for ra in m["ras"]:
        hdr(ws, 4, col, f"RA{ra['n']}", 8, SUB, Font(name="Calibri", size=10, bold=True, color="1F4E79"))
        w = ws.cell(3, col, ra["peso_ra"]); w.fill, w.border, w.alignment = SUB, BD, CENTER
        ra_cols[ra["n"]] = col
        col += 1
    hdr(ws, 4, col, "CE pendientes", 10); pend_col = col; col += 1
    hdr(ws, 4, col, "Nota módulo", 10, HF); nm_col = col; col += 1
    hdr(ws, 4, col, "Calificación", 10, HF); cal_col = col
    ws.row_dimensions[4].height = 32
    ra_first, ra_last = L(min(ra_cols.values())), L(max(ra_cols.values()))
    ce_first, ce_last = L(min(ce_cols.values())), L(max(ce_cols.values()))
    # evidencias por CE
    cov = {k: [] for k in ces}
    for s in m["sdas"]:
        for e in s["evidencias"]:
            for k in e["ces"]:
                cov[k].append((ev_cols[(s["n"], e["id"])], e["peso"]))
    for r in range(5, 5 + n):
        ws.cell(r, 1, name_ref(r)).border = BD
        for c in ev_cols.values():
            ws.cell(r, c).fill, ws.cell(r, c).border = INPUT, BD
            dv.add(ws.cell(r, c))
        for sn, a, b in sda_blocks:
            c = ws.cell(r, sda_nota[sn], f'=IF(COUNT({L(a)}{r}:{L(b)}{r})={b-a+1},SUMPRODUCT({L(a)}{r}:{L(b)}{r},{L(a)}$3:{L(b)}$3)/100,"")')
            c.fill, c.border, c.number_format = CALC, BD, "0.00"
        for k, c3 in ce_cols.items():
            if k in dual and k in feoe_refs:
                fc = f"='FEOE seguimiento'!{L(feoe_refs[k])}{r}"
            else:
                lst = cov[k]
                cnt = f"COUNT({','.join(f'{L(c)}{r}' for c, _ in lst)})"
                num = "+".join(f"{w}*{L(c)}{r}" for c, w in lst)
                den = sum(w for _, w in lst)
                fc = f'=IF({cnt}={len(lst)},({num})/{den},"")'
            c = ws.cell(r, c3, fc)
            c.fill, c.border, c.number_format = CALC, BD, "0.0"
        for raN, (a, b) in ra_ranges.items():
            c = ws.cell(r, ra_cols[raN], f'=IF(COUNT({L(a)}{r}:{L(b)}{r})={b-a+1},SUMPRODUCT({L(a)}{r}:{L(b)}{r},{L(a)}$3:{L(b)}$3)/100,"")')
            c.fill, c.border, c.number_format = CALC, BD, "0.00"
        c = ws.cell(r, pend_col, f"={len(ces)}-COUNT({ce_first}{r}:{ce_last}{r})"); c.fill, c.border = CALC, BD
        c = ws.cell(r, nm_col, f'=IF(COUNT({ra_first}{r}:{ra_last}{r})={len(ra_cols)},SUMPRODUCT({ra_first}{r}:{ra_last}{r},{ra_first}$3:{ra_last}$3)/100,"")')
        c.fill, c.border, c.number_format = CALC, BD, "0.00"
        c = ws.cell(r, cal_col, f'=IF({L(nm_col)}{r}="","",ROUND({L(nm_col)}{r},0))'); c.fill, c.border = CALC, BD
        c.font = Font(bold=True)
    ws.freeze_panes = "B5"

def instrucciones(wb, m):
    ws = wb.create_sheet("Cómo usarlo", 0)
    ws.column_dimensions["A"].width = 120
    lines = [
        ("Cuaderno del profesor · cómo usarlo", True),
        ("1. Escriba los nombres del grupo en la hoja «Alumnado». El resto de hojas los toman de ahí.", False),
        ("2. Hojas por evaluación: una fila por sesión real (horario × calendario − días no lectivos), con la planificación, los CE, la evidencia que se entrega y dos columnas libres para el diario de aula y las incidencias.", False),
        ("3. «Asistencia» y «Actitud»: registro por sesión y por evaluación. La actitud profesional no modifica la nota de los CE.", False),
        ("4. «Calificaciones»: escriba las notas de cada evidencia en las casillas amarillas. Las notas de SdA, CE, RA y del módulo se calculan solas con los pesos de la Tabla maestra.", False),
        ("5. «FEOE seguimiento»: por cada CE dualizado anote la nota del tutor de empresa o, si la empresa no lo cubre, la de la alternativa en el centro. Esa nota alimenta la hoja «Calificaciones».", False),
        ("6. Recuperación: se escribe la nueva nota encima de la anterior (la recuperada sustituye) y se deja constancia en el diario de aula con la fecha.", False),
        ("7. Las casillas azules son fórmulas: no se tocan. Si no hay todos los datos, quedan vacías.", False),
        ("8. Este cuaderno no contiene datos personales. No anote datos de necesidades específicas de apoyo sin acreditación oficial.", False),
    ]
    for i, (t, b) in enumerate(lines, 1):
        c = ws.cell(i, 1, t)
        c.font = Font(name="Calibri", size=13 if b else 11, bold=b, color="1F4E79" if b else "000000")
        c.alignment = Alignment(wrap_text=True, vertical="top")

def add_extra_sheets(wb, m):
    alumnado(wb, m)
    asistencia(wb, m)
    actitud(wb, m)
    refs = feoe(wb, m)
    calificaciones(wb, m, refs)
    instrucciones(wb, m)
