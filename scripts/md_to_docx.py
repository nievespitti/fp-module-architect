#!/usr/bin/env python3
"""Convierte un .md a .docx (títulos, listas, tablas, negrita, cursiva). Quita comentarios HTML.

Uso: python md_to_docx.py <entrada.md> [salida.docx]
"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INLINE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")

def runs(par, text, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1]); r.font.name = "Consolas"
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            r = par.add_run(part[1:-1]); r.italic = True
        else:
            r = par.add_run(part)
        if bold:
            r.bold = True

def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def table(doc, rows):
    n = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=n)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j in range(n):
            c = t.cell(i, j)
            c.text = ""
            p = c.paragraphs[0]
            runs(p, r[j] if j < len(r) else "", bold=(i == 0))
            for run in p.runs:
                run.font.size = Pt(9)
                if i == 0:
                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if i == 0:
                shade(c, "1F4E79")
    doc.add_paragraph()

def convert(src, dst):
    text = re.sub(r"<!--.*?-->", "", open(src, encoding="utf-8").read(), flags=re.S)
    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2); s.top_margin = s.bottom_margin = Cm(2)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    lines, i = text.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-|]+\|$", lines[i + 1].strip()):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|[\s:\-|]+\|$", lines[i].strip()):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
                i += 1
            table(doc, rows)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)", ln)
        if m:
            doc.add_heading(m.group(2).strip(), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", ln):
            runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", ln))
        elif re.match(r"^\s*\d+[.)]\s+", ln):
            runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+[.)]\s+", "", ln))
        elif ln.strip() in ("---", "***"):
            pass
        elif ln.strip():
            runs(doc.add_paragraph(), ln.strip().lstrip("> "))
        i += 1
    doc.save(dst)

if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    src = sys.argv[1]
    convert(src, sys.argv[2] if len(sys.argv) > 2 else re.sub(r"\.md$", ".docx", src))
    print("OK")
