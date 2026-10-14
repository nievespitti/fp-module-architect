import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()
wb.remove(wb.active)  # remove default sheet

# Professional palette
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Calibri", size=13, bold=True, color="1F4E79")
SUBTITLE_FONT = Font(name="Calibri", size=10, italic=True, color="595959")
CALENDAR_INFO_FONT = Font(name="Calibri", size=9.5, bold=True, color="C00000")

ZEBRA_FILL = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

# Highlight styles
EVAL_FILL = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")     # Blue soft
EVAL_FONT = Font(name="Calibri", size=10, bold=True, color="1F4E79")

FESTIVO_FILL = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # Coral/Orange soft
FESTIVO_FONT = Font(name="Calibri", size=10, bold=True, color="C65911")

ENTREGA_FILL = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")  # Yellow soft
ENTREGA_FONT = Font(name="Calibri", size=10, bold=True, color="806000")

THIN_BORDER = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)
HEADER_BORDER = Border(
    left=Side(style='thin', color='1F4E79'),
    right=Side(style='thin', color='1F4E79'),
    top=Side(style='thin', color='1F4E79'),
    bottom=Side(style='medium', color='1F4E79')
)

COLUMNS = [
    ("Sesión", 8, "center"),
    ("Fecha", 14, "center"),
    ("Día", 11, "center"),
    ("Horas", 7, "center"),
    ("SdA / Unidad", 22, "left"),
    ("RA / CE", 12, "center"),
    ("Planificación / Contenido Previsto", 38, "left"),
    ("Diario de Aula / Trabajo Realizado (Registro Profesor)", 40, "left"),
    ("Incidencias / Tareas / Observaciones", 30, "left"),
    ("Evidencia Alexia", 16, "center")
]

def format_sheet(ws, title, subtitle, calendar_note):
    ws.merge_cells("A1:J1")
    ws["A1"] = title
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    
    ws.merge_cells("A2:J2")
    ws["A2"] = subtitle
    ws["A2"].font = SUBTITLE_FONT
    ws["A2"].alignment = Alignment(horizontal="left", vertical="center")

    ws.merge_cells("A3:J3")
    ws["A3"] = calendar_note
    ws["A3"].font = CALENDAR_INFO_FONT
    ws["A3"].alignment = Alignment(horizontal="left", vertical="center")
    
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 18
    
    # Headers on row 4
    ws.row_dimensions[4].height = 28
    for col_idx, (col_name, col_width, col_align) in enumerate(COLUMNS, 1):
        cell = ws.cell(row=4, column=col_idx, value=col_name)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = HEADER_BORDER
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

def populate_rows(ws, sessions_data):
    current_row = 5
    for row_data in sessions_data:
        ws.row_dimensions[current_row].height = 26 if len(str(row_data[6])) < 40 else 36
        row_type = row_data[10] if len(row_data) > 10 else "NORMAL"
        
        if row_type == "FESTIVO":
            row_fill = FESTIVO_FILL
            row_font = FESTIVO_FONT
        elif row_type == "EVALUACION":
            row_fill = EVAL_FILL
            row_font = EVAL_FONT
        elif row_type == "ENTREGA":
            row_fill = ENTREGA_FILL
            row_font = ENTREGA_FONT
        else:
            is_zebra = (current_row % 2 == 0)
            row_fill = ZEBRA_FILL if is_zebra else WHITE_FILL
            row_font = Font(name="Calibri", size=10, bold=False)

        for col_idx in range(10):
            val = row_data[col_idx]
            cell = ws.cell(row=current_row, column=col_idx+1, value=val)
            cell.fill = row_fill
            cell.border = THIN_BORDER
            align_type = COLUMNS[col_idx][2]
            cell.alignment = Alignment(horizontal=align_type, vertical="center", wrap_text=True)
            cell.font = row_font if (row_type in ["FESTIVO", "EVALUACION", "ENTREGA"] and col_idx in [0, 1, 4, 6, 9]) else Font(name="Calibri", size=10)
        current_row += 1

# ==================== 1ª EVALUACIÓN ====================
# Calendario v5: Inicio clases: X-16-9-2026. Eval Inicial: M-20-10-2026. 1ª Eval: M-15-12-2026. Vacaciones Navidad: X-23-12-2026 al V-8-1-2027.
# Días libre disposición / no lectivos: L-02-11-2026, V-04-12-2026, L-07-12-2026, M-08-12-2026 (Festivo Inmaculada).
s1_data = [
    # Semana 0: 16-17 sep
    (1, "16/09/2026", "Miércoles", 2, "Semana 0", "-", "INICIO DE CLASES. Presentación módulo, normas taller, NUSS y Cl@ve", "", "", "-", "NORMAL"),
    (2, "17/09/2026", "Jueves", 2, "Semana 0", "-", "Estructura curricular FP, criterios de evaluación y cuaderno Alexia", "", "", "-", "NORMAL"),
    # Semana 1: 22-24 sep
    (3, "22/09/2026", "Martes", 2, "SdA1 · Redes", "RA1", "Act. 1: ¿Qué es una red local? Escenario pyme inicial", "", "", "-", "NORMAL"),
    (4, "23/09/2026", "Miércoles", 2, "SdA1 · Redes", "RA1", "Componentes de una red: clientes, servidores, medios y servicios", "", "", "-", "NORMAL"),
    (5, "24/09/2026", "Jueves", 2, "SdA1 · Redes", "RA1", "Act. 2: Introducción a Cisco Packet Tracer. Interfaz y modos", "", "", "-", "NORMAL"),
    # Semana 2: 29 sep - 1 oct
    (6, "29/09/2026", "Martes", 2, "SdA1 · Redes", "RA1.a, c", "Act. 2: Primera topología PC0-Switch-PC1. Conexión de cable directo", "", "", "-", "NORMAL"),
    (7, "30/09/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.a, RA4", "Configuración IPv4 estática (/24) y máscara de subred", "", "", "-", "NORMAL"),
    (8, "01/10/2026", "Jueves", 2, "SdA1 · Redes", "RA1.a, RA4", "Comprobación de conectividad con comando ping. Entrega EV1", "", "", "S1_EV1", "ENTREGA"),
    # Semana 3: 6-8 oct
    (9, "06/10/2026", "Martes", 2, "SdA1 · Redes", "RA1.a, RA4", "Direccionamiento IPv4: porción de red vs. porción de host", "", "", "-", "NORMAL"),
    (10, "07/10/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.a, c", "Act. 3: Diagnóstico ¿Por qué no funciona? Fallos provocados", "", "", "-", "NORMAL"),
    (11, "08/10/2026", "Jueves", 2, "SdA1 · Redes", "RA1.a, c", "Resolución de incidencias IP y cableado. Entrega ficha y PKT", "", "", "S1_EV2", "ENTREGA"),
    # Semana 4: 13-15 oct (12 oct Fiesta Nacional - Lunes)
    (12, "13/10/2026", "Martes", 2, "SdA1 · Redes", "RA1.c", "Act. 4: Dispositivos de interconexión (NIC, Hub, Switch, AP)", "", "", "-", "NORMAL"),
    (13, "14/10/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.c", "Funciones del Router y concepto de Gateway (puerta de enlace)", "", "", "-", "NORMAL"),
    (14, "15/10/2026", "Jueves", 2, "SdA1 · Redes", "RA1.c", "Matriz comparativa de dispositivos y casos prácticos de uso", "", "", "-", "NORMAL"),
    # Semana 5: 20-22 oct (20/10 EVALUACIÓN INICIAL oficial v5)
    (15, "20/10/2026", "Martes", 2, "EVALUACIÓN INICIAL", "RA1", "SESIÓN OFICIAL EVALUACIÓN INICIAL (Calendario v5). Chequeo diagnóstico", "", "Informe tutoría", "Eval. Inicial", "EVALUACION"),
    (16, "21/10/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.a", "Act. 5: Modelos de comunicación OSI y TCP/IP", "", "", "-", "NORMAL"),
    (17, "22/10/2026", "Jueves", 2, "SdA1 · Redes", "RA1.a", "Packet Tracer Simulation Mode: Inspección de PDUs, ARP e ICMP", "", "", "-", "NORMAL"),
    # Semana 6: 27-29 oct
    (18, "27/10/2026", "Martes", 2, "SdA1 · Redes", "RA1.b", "Act. 6: Clasificación de redes según alcance (PAN, LAN, MAN, WAN)", "", "", "-", "NORMAL"),
    (19, "28/10/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.d", "Act. 7: Medios de transmisión: Cobre (UTP/STP), Fibra y Wi-Fi", "", "", "-", "NORMAL"),
    (20, "29/10/2026", "Jueves", 2, "SdA1 · Redes", "RA1.g, h", "Act. 8: Topologías físicas y lógicas. Entrega bloque fundamentos", "", "", "S1_EV3", "ENTREGA"),
    # Semana 7: 3-5 nov (02/11 Lunes no lectivo)
    (21, "03/11/2026", "Martes", 2, "SdA1 · Redes", "RA1.e", "Act. 9: Interpretación de planos y mapas físicos de telecomunicación", "", "", "-", "NORMAL"),
    (22, "04/11/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.f", "Herramientas de dibujo digital: Draw.io / Excalidraw", "", "", "-", "NORMAL"),
    (23, "05/11/2026", "Jueves", 2, "SdA1 · Redes", "RA1.e, f", "Diseño de plano con simbología normalizada. Entrega plano digital", "", "", "S1_EV4", "ENTREGA"),
    # Semana 8: 10-12 nov
    (24, "10/11/2026", "Martes", 2, "SdA1 · Redes", "RA1.a-h", "Act. 10: Lanzamiento Proyecto Final SdA1 (Red Pequeña Empresa)", "", "", "-", "NORMAL"),
    (25, "11/11/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.a-h", "Trabajo en proyecto: dimensionamiento, direccionamiento y plano", "", "", "-", "NORMAL"),
    (26, "12/11/2026", "Jueves", 2, "SdA1 · Redes", "RA1.a-h", "Finalización de maqueta PKT y memoria técnica pyme", "", "", "S1_EV5", "ENTREGA"),
    # Semana 9: 17-19 nov
    (27, "17/11/2026", "Martes", 2, "SdA1 · Redes", "RA1.a-h", "Act. 11: Prueba Práctica Individual de Demostración (Grupo A/B)", "", "", "S1_EV6", "ENTREGA"),
    (28, "18/11/2026", "Miércoles", 2, "SdA1 · Redes", "RA1.a-h", "Prueba Práctica Individual de Demostración (Grupo B/A) y defensas", "", "", "S1_EV6", "ENTREGA"),
    (29, "19/11/2026", "Jueves", 2, "SdA1 · Redes", "RA1", "Feedback individual SdA1, cierre de notas Alexia y recuperación", "", "", "Cierre SdA1", "NORMAL"),
    # Semana 10: 24-26 nov (Inicio SdA2 · Del plano al cable)
    (30, "24/11/2026", "Martes", 2, "SdA2 · Cableado", "RA2.a, c", "Lanzamiento SdA2: Estándares de cableado estructurado (EIA/TIA 568)", "", "", "-", "NORMAL"),
    (31, "25/11/2026", "Miércoles", 2, "SdA2 · Cableado", "RA2.c, d", "Cables de cobre: categorías (Cat5e, Cat6, Cat6A), blindajes (UTP/FTP)", "", "", "-", "NORMAL"),
    (32, "26/11/2026", "Jueves", 2, "SdA2 · Cableado", "RA2.d", "Canalizaciones, tubos, bandejas y radios de curvatura", "", "", "-", "NORMAL"),
    # Semana 11: 1-3 dic (04/12 Viernes libre disposición)
    (33, "01/12/2026", "Martes", 2, "SdA2 · Cableado", "RA2.c, g", "Herramientas de taller: crimpadora, pelacables, impactadora (Krone)", "", "", "-", "NORMAL"),
    (34, "02/12/2026", "Miércoles", 2, "SdA2 · Cableado", "RA2.g, 6.b", "Práctica de crimpado RJ-45 directo (norma T568B). Seguridad taller", "", "", "S2_EV1", "ENTREGA"),
    (35, "03/12/2026", "Jueves", 2, "SdA2 · Cableado", "RA2.h", "Comprobación con tester de cables (mapa de hilos, continuidad)", "", "", "S2_EV2", "ENTREGA"),
    # Semana 12: 8-10 dic (07/12 Lunes libre disp., 08/12 Martes Festivo Inmaculada)
    (36, "08/12/2026", "Martes", 0, "FESTIVO NACIONAL", "-", "FESTIVO NACIONAL: Inmaculada Concepción (Día no lectivo)", "-", "-", "-", "FESTIVO"),
    (37, "09/12/2026", "Miércoles", 2, "SdA2 · Cableado", "RA2.f, g", "Armarios rack de comunicaciones: bastidores, paneles y guías", "", "", "-", "NORMAL"),
    (38, "10/12/2026", "Jueves", 2, "SdA2 · Cableado", "RA2.g, i", "Montaje en rosetas hembra RJ-45 (keystone) y patch panel", "", "", "-", "NORMAL"),
    # Semana 13: 15-17 dic (15/12 PRIMERA EVALUACIÓN oficial v5)
    (39, "15/12/2026", "Martes", 2, "1ª EVALUACIÓN", "RA1, RA2", "SESIÓN DE 1.ª EVALUACIÓN (Calendario v5). Cierre de notas en Alexia", "", "Boletines 1ª Eval", "1ª EVAL", "EVALUACION"),
    (40, "16/12/2026", "Miércoles", 2, "SdA2 · Cableado", "RA2.g, h, i", "Taller: Cableado rack-roseta, conexionado y etiquetado", "", "", "S2_EV3", "ENTREGA"),
    (41, "17/12/2026", "Jueves", 2, "SdA2 · Cableado", "RA2.h, j", "Certificación básica de enlaces con comprobador digital", "", "", "-", "NORMAL"),
    # Semana 14: 22 dic (Último día lectivo previo a vacaciones Navidad X-23-12 al V-08-01)
    (42, "22/12/2026", "Martes", 2, "Cierre Trimestre", "RA1, RA2", "Entrega de boletines de 1ª evaluación, balance trimestral y recuperación", "", "", "Cierre T1", "NORMAL")
]

# ==================== 2ª EVALUACIÓN ====================
# Calendario v5: FEOE: Del M-2-2-2027 al V-5-3-2027. Segunda evaluación: M-9-3-2027.
# Libre disposición: L-01-02-2027 y L-08-02-2027 (Carnaval).
s2_data = [
    # Semana 15: 12-14 ene (Vuelta de Navidad el 11/01)
    (43, "12/01/2027", "Martes", 2, "SdA2 · Cableado", "RA2.d-j", "Proyecto final de instalación de cableado en maqueta de taller", "", "", "-", "NORMAL"),
    (44, "13/01/2027", "Miércoles", 2, "SdA2 · Cableado", "RA2.d-j", "Finalización de montaje, peinado de rack y etiquetado normalizado", "", "", "S2_EV4", "ENTREGA"),
    (45, "14/01/2027", "Jueves", 2, "SdA2 · Cableado", "RA2.a-j", "Prueba práctica individual de confección y testeo de enlaces", "", "", "S2_EV5", "ENTREGA"),
    # Semana 16: 19-21 ene (SdA3 · Conectamos la red)
    (46, "19/01/2027", "Martes", 2, "SdA3 · Interconexión", "RA3.a, b", "Lanzamiento SdA3: Adaptadores de red (NIC Ethernet, PCIe, USB, Wi-Fi)", "", "", "-", "NORMAL"),
    (47, "20/01/2027", "Miércoles", 2, "SdA3 · Interconexión", "RA3.b, c", "Instalación de adaptadores en equipos reales y configuración física", "", "", "-", "NORMAL"),
    (48, "21/01/2027", "Jueves", 2, "SdA3 · Interconexión", "RA3.d, e", "Montaje físico del Switch en rack y conexión a patch panels", "", "", "S3_EV1", "ENTREGA"),
    # Semana 17: 26-28 ene (SdA4 · PRL y Seguridad previo a FEOE)
    (49, "26/01/2027", "Martes", 2, "SdA4 · PRL / RA6", "RA6.a, b, d", "Lanzamiento SdA4: Auditoría de seguridad y riesgos en puesto TIC", "", "", "-", "NORMAL"),
    (50, "27/01/2027", "Miércoles", 2, "SdA4 · PRL / RA6", "RA6.d, e, g", "EPIs, ergonomía, gestión selectiva de residuos informáticos (RAEE)", "", "", "S4_EV1", "ENTREGA"),
    (51, "28/01/2027", "Jueves", 2, "SdA4 · PRL / RA6", "RA6.a-h", "Cierre formal RA6 previo a FEOE: Demostración y checklist de taller", "", "", "S4_EV2", "ENTREGA"),
    # Lunes 01/02/2027 libre disposición
    # Periodo FEOE oficial v5: Del M-02-02-2027 al V-05-03-2027
    (52, "02/02/2027", "Martes", 2, "FEOE / Empresa", "FCT / Dual", "INICIO OFICIAL PERIODO FEOE (Calendario v5). Estancia en empresa", "", "Empresa", "-", "EVALUACION"),
    (53, "08/02/2027", "Lunes", 0, "LIBRE DISPOSICIÓN", "-", "DÍA DE LIBRE DISPOSICIÓN (Carnaval) según Calendario del Centro", "-", "-", "-", "FESTIVO"),
    (54, "15/02/2027", "Lunes", 2, "FEOE / Tutoría", "Seguimiento", "Tutoría de seguimiento FEOE en centro educativo: revisión de tareas", "", "Cuaderno FEOE", "-", "NORMAL"),
    (55, "22/02/2027", "Lunes", 2, "FEOE / Tutoría", "Seguimiento", "Tutoría de seguimiento FEOE: resolución de dudas e incidencias técnicas", "", "Cuaderno FEOE", "-", "NORMAL"),
    (56, "01/03/2027", "Lunes", 2, "FEOE / Tutoría", "Seguimiento", "Tutoría final de empresa: recopilación de valoraciones y firmas", "", "Cuaderno FEOE", "-", "NORMAL"),
    (57, "05/03/2027", "Viernes", 2, "FEOE / Cierre", "FCT / Dual", "FINALIZACIÓN OFICIAL PERIODO FEOE (Calendario v5)", "", "Acta FEOE", "-", "EVALUACION"),
    # Semana 18: 09-11 mar (09/03 SEGUNDA EVALUACIÓN oficial v5)
    (58, "09/03/2027", "Martes", 2, "2ª EVALUACIÓN", "RA2, RA3, RA6", "SESIÓN DE 2.ª EVALUACIÓN (Calendario v5). Cierre de notas en Alexia", "", "Boletines 2ª Eval", "2ª EVAL", "EVALUACION"),
    (59, "10/03/2027", "Miércoles", 2, "SdA3 · Interconexión", "RA3.f, g", "Puesta en común experiencias FEOE y comprobación de enlaces de red", "", "", "-", "NORMAL"),
    (60, "11/03/2027", "Jueves", 2, "SdA3 · Interconexión", "RA3.f, g", "Verificación lógica (LEDs, ping, diagnóstico) y prueba práctica", "", "", "S3_EV2", "ENTREGA")
]

# ==================== 3ª EVALUACIÓN ====================
# Calendario v5: Semana Santa: L-22-3-2027 al V-26-3-2027. Libre disp: L-31-5-2027.
# Final clases: X-2-6-2027. Semana exámenes: J-3-6-2027 al M-8-6-2027.
# Evaluación final: J-10-6-2027. Entrega notas y fin de curso: M-22-6-2027.
s3_data = [
    # Semana 19: 16-18 mar (Lanzamiento SdA5 · Configuramos la red)
    (61, "16/03/2027", "Martes", 2, "SdA5 · Configuración", "RA4.a, b", "Lanzamiento SdA5: Arquitectura de redes inalámbricas WLAN (802.11)", "", "", "-", "NORMAL"),
    (62, "17/03/2027", "Miércoles", 2, "SdA5 · Configuración", "RA4.c, d", "Configuración de Puntos de Acceso (AP): SSIDs, canales y frecuencias", "", "", "-", "NORMAL"),
    (63, "18/03/2027", "Jueves", 2, "SdA5 · Configuración", "RA4.d, e", "Práctica de conexión cliente Wi-Fi y comprobación de cobertura", "", "", "S5_EV1", "ENTREGA"),
    # Semana Santa: L-22-3-2027 al V-26-3-2027 (Vacaciones)
    # Semana 20: 30 mar - 1 abr
    (64, "30/03/2027", "Martes", 2, "SdA5 · Configuración", "RA4.g, h", "Protocolo IP en profundidad: clases, máscaras, puertas de enlace", "", "", "-", "NORMAL"),
    (65, "31/03/2027", "Miércoles", 2, "SdA5 · Configuración", "RA4.h", "Subnetting básico /24, /25, /26 en Packet Tracer", "", "", "-", "NORMAL"),
    (66, "01/04/2027", "Jueves", 2, "SdA5 · Configuración", "RA4.h", "Práctica guiada de asignación de direcciones estáticas y resolución", "", "", "-", "NORMAL"),
    # Semana 21: 6-8 abr
    (67, "06/04/2027", "Martes", 2, "SdA5 · Configuración", "RA4.i", "Seguridad básica en redes locales: WPA2/WPA3 Personal, filtrado MAC", "", "", "-", "NORMAL"),
    (68, "07/04/2027", "Miércoles", 2, "SdA5 · Configuración", "RA4.i", "Aislamiento de clientes y buenas prácticas de protección de red", "", "", "S5_EV2", "ENTREGA"),
    (69, "08/04/2027", "Jueves", 2, "SdA5 · Configuración", "RA4.j", "Concepto de Redes de Área Local Virtuales (VLAN). Beneficios", "", "", "-", "NORMAL"),
    # Semana 22: 13-15 abr
    (70, "13/04/2027", "Martes", 2, "SdA5 · Configuración", "RA4.j", "Configuración de VLANs de acceso en switch Cisco (CLI y Packet Tracer)", "", "", "-", "NORMAL"),
    (71, "14/04/2027", "Miércoles", 2, "SdA5 · Configuración", "RA4.j", "Configuración de enlaces troncales (Trunk 802.1Q)", "", "", "-", "NORMAL"),
    (72, "15/04/2027", "Jueves", 2, "SdA5 · Configuración", "RA4.j", "Reto VLAN: Aislamiento por departamentos en Packet Tracer", "", "", "S5_EV3", "ENTREGA"),
    # Semana 23: 20-22 abr
    (73, "20/04/2027", "Martes", 2, "SdA5 · Configuración", "RA4.a-j", "Proyecto TecnoCanarias: Integración WLAN + VLAN + Seguridad", "", "", "-", "NORMAL"),
    (74, "21/04/2027", "Miércoles", 2, "SdA5 · Configuración", "RA4.a-j", "Finalización y entrega de maqueta de red configurada", "", "", "S5_EV4", "ENTREGA"),
    (75, "22/04/2027", "Jueves", 2, "SdA5 · Configuración", "RA4.a-j", "Prueba práctica individual en aula de configuración de red y VLAN", "", "", "S5_EV5", "ENTREGA"),
    # Semana 24: 27-29 abr (Lanzamiento SdA6 · La red no funciona)
    (76, "27/04/2027", "Martes", 2, "SdA6 · Incidencias", "RA5.a, b", "Lanzamiento SdA6: Metodología de resolución de incidencias (Troubleshooting)", "", "", "-", "NORMAL"),
    (77, "28/04/2027", "Miércoles", 2, "SdA6 · Incidencias", "RA5.a, b", "Gestión de tickets de soporte: tipificación de fallos hardware vs software", "", "", "-", "NORMAL"),
    (78, "29/04/2027", "Jueves", 2, "SdA6 · Incidencias", "RA5.c, d", "Comandos de diagnóstico de red: ping, tracert/traceroute, netstat, arp", "", "", "S6_EV1", "ENTREGA"),
    # Semana 25: 4-6 may (01/05 Sábado Fiesta del Trabajo)
    (79, "04/05/2027", "Martes", 2, "SdA6 · Incidencias", "RA5.d, e", "Análisis de tráfico con Wireshark y Packet Tracer Simulation", "", "", "-", "NORMAL"),
    (80, "05/05/2027", "Miércoles", 2, "SdA6 · Incidencias", "RA5.e, f", "Taller práctico: localización de averías físicas (latiguillo cortado, puerto dañado)", "", "", "-", "NORMAL"),
    (81, "06/05/2027", "Jueves", 2, "SdA6 · Incidencias", "RA5.f, g", "Sustitución de equipamiento averiado y restauración de configuraciones", "", "", "S6_EV2", "ENTREGA"),
    # Semana 26: 11-13 may
    (82, "11/05/2027", "Martes", 2, "SdA6 · Incidencias", "RA5.g, h", "Resolución de incidencias lógicas complejas (duplicidad IP, máscara errónea)", "", "", "-", "NORMAL"),
    (83, "12/05/2027", "Miércoles", 2, "SdA6 · Incidencias", "RA5.h", "Elaboración de informes de asistencia técnica y partes de trabajo", "", "", "S6_EV3", "ENTREGA"),
    (84, "13/05/2027", "Jueves", 2, "SdA6 · Incidencias", "RA5.a-h", "Reto integrador de incidencias: maratón de resolución por equipos", "", "", "-", "NORMAL"),
    # Semana 27: 18-20 may
    (85, "18/05/2027", "Martes", 2, "SdA6 · Incidencias", "RA5.a-h", "Prueba práctica individual de diagnóstico y reparación de averías", "", "", "S6_EV4", "ENTREGA"),
    (86, "19/05/2027", "Miércoles", 2, "SdA6 · Incidencias", "RA5", "Cierre de SdA6, revisión de partes y actualización de notas Alexia", "", "", "Cierre SdA6", "NORMAL"),
    (87, "20/05/2027", "Jueves", 2, "Repaso / Recuperación", "RA1 - RA6", "Taller de recuperación de evidencias prácticas pendientes (RA1, RA2, RA3)", "", "", "-", "NORMAL"),
    # Semana 28: 25-27 may (30/05 Domingo Día de Canarias, 31/05 Lunes Libre Disposición)
    (88, "25/05/2027", "Martes", 2, "Repaso / Recuperación", "RA1 - RA6", "Taller de recuperación de evidencias prácticas pendientes (RA4, RA5, RA6)", "", "", "-", "NORMAL"),
    (89, "26/05/2027", "Miércoles", 2, "Repaso / Recuperación", "Todos", "Simulacro global de prueba práctica final", "", "", "-", "NORMAL"),
    (90, "27/05/2027", "Jueves", 2, "Cierre Clases", "Todos", "Balance de aprendizajes del módulo y preparación para examen ordinario", "", "", "-", "NORMAL"),
    # Semana 29: 1-3 jun (31/05 Lunes no lectivo; 02/06 Miércoles FINAL CLASES; 03-08 jun EXÁMENES)
    (91, "01/06/2027", "Martes", 2, "Cierre Ordinario", "Todos", "Última clase ordinaria de aula: resolución de dudas previas a examen", "", "", "-", "NORMAL"),
    (92, "02/06/2027", "Miércoles", 2, "FINAL CLASES ORDINARIAS", "Todos", "FINAL DE CLASES ORDINARIAS (Calendario v5). Cierre de firmas de asistencia", "", "", "Fin Clases", "EVALUACION"),
    (93, "03/06/2027", "Jueves", 2, "SEMANA DE EXÁMENES", "RA1 - RA6", "SEMANA DE EXÁMENES (Calendario v5: 03-08 jun). Prueba global ordinaria", "", "Prueba global", "Examen Final", "EVALUACION"),
    # Semana 30: 10 jun (J-10-06-2027 EVALUACIÓN FINAL) y 22 jun (M-22-06-2027 ENTREGA NOTAS)
    (94, "10/06/2027", "Jueves", 2, "EVALUACIÓN FINAL", "RA1 - RA6", "SESIÓN OFICIAL DE EVALUACIÓN FINAL (Calendario v5). Cierre definitivo Alexia", "", "Actas Oficiales", "EVAL FINAL", "EVALUACION"),
    (95, "22/06/2027", "Martes", 2, "ENTREGA DE NOTAS", "Todos", "ENTREGA DE NOTAS Y FIN DEL CURSO (Calendario v5)", "", "Boletines Finales", "Fin Curso", "EVALUACION")
]

# Create Sheets
ws1 = wb.create_sheet("1ª Evaluación")
format_sheet(ws1, "Cuaderno del Profesor · Diario de Aula · 1.ª Evaluación (2026-2027)", 
             "Módulo: 0225 Redes Locales · 1º SMR (6h/sem: Martes, Miércoles, Jueves) · SdA1 (RA1) e inicio SdA2 (RA2)",
             "Hitos Oficiales v5: Inicio: 16/09/2026 | Eval. Inicial: 20/10/2026 | 1.ª Eval: 15/12/2026 | No lectivos: L-02/11, V-04/12, L-07/12, M-08/12 (Festivo) | Navidad: 23/12/26 al 08/01/27")
populate_rows(ws1, s1_data)

ws2 = wb.create_sheet("2ª Evaluación")
format_sheet(ws2, "Cuaderno del Profesor · Diario de Aula · 2.ª Evaluación (2026-2027)", 
             "Módulo: 0225 Redes Locales · 1º SMR · SdA2 (RA2), SdA3 (RA3), SdA4 (RA6) y Periodo FEOE (Formación en Empresa)",
             "Hitos Oficiales v5: FEOE: 02/02/2027 al 05/03/2027 | 2.ª Eval: 09/03/2027 | Libre disposición: L-01/02 y L-08/02 (Carnaval)")
populate_rows(ws2, s2_data)

ws3 = wb.create_sheet("3ª Evaluación")
format_sheet(ws3, "Cuaderno del Profesor · Diario de Aula · 3.ª Evaluación (2026-2027)", 
             "Módulo: 0225 Redes Locales · 1º SMR · SdA5 (RA4), SdA6 (RA5) y Evaluación Final Ordinaria",
             "Hitos Oficiales v5: Semana Santa: 22 al 26/03/2027 | Libre disp: L-31/05/2027 | Fin Clases: 02/06/2027 | Exámenes: 03 al 08/06/2027 | Eval. Final: 10/06/2027 | Notas: 22/06/2027")
populate_rows(ws3, s3_data)

output_path = "Cuaderno_Profesor_Diario_Aula_2026_2027.xlsx"
wb.save(output_path)
print(f"Regenerado con exito conforme al Calendario v5: {output_path}")
